"""Run: python -m unittest discover -s tests   (no audio hardware / Whisper needed)"""
import json
import sys
import tempfile
import threading
import unittest
from datetime import datetime
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from docx import Document  # noqa: E402

from mom_recorder import mom_generator as mg  # noqa: E402
from mom_recorder.audio import Segmenter, to_16k  # noqa: E402
from mom_recorder.config import DEFAULTS  # noqa: E402
from mom_recorder.transcript_doc import TranscriptWriter  # noqa: E402

REPLY = {
    "overview": "The meeting reviewed the Q4 release plan.",
    "scope": "To agree release scope and owners.",
    "mode": "Online - Microsoft Teams",
    "external_organization": "Acme",
    "external_participants": ["Jane Doe", "John Roe"],
    "internal_participants": ["Kishore"],
    "discussion_points": [{"title": "Release Scope", "details": "Scope was frozen."},
                          {"title": "Testing & QA", "details": "Automated tests will run nightly."},
                          {"title": "Hosting", "details": "Cloud hosting preferred."}],
    "action_items": [{"activity": "Share test plan", "organization": "Mahat", "owner": "Kishore",
                      "due_by": "20 Oct 2026", "comments": ""},
                     {"activity": "Confirm hosting", "organization": "Acme", "owner": "Jane Doe",
                      "due_by": "", "comments": ""},
                     {"activity": "Send VAPT docs", "organization": "Acme", "owner": "", "due_by": "", "comments": ""}],
}
SEEN = []


class H(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        SEEN.append(body)
        out = json.dumps({"response": json.dumps(REPLY) if body.get("format") == "json" else "- note"}).encode()
        self.send_response(200)
        self.end_headers()
        self.wfile.write(out)

    def do_GET(self):
        self.send_response(200)
        self.end_headers()
        self.wfile.write(json.dumps({"models": [{"name": "llama3"}]}).encode())


class Tests(unittest.TestCase):
    def test_segmenter_and_resample(self):
        sr = 48000
        seg = Segmenter("x", sr, 0.01, 0.5, 25)
        tone = (0.3 * np.sin(np.arange(sr) * 0.05)).astype(np.float32)
        quiet = np.zeros(1024, dtype=np.float32)
        out, t = [], 0.0
        for b in [quiet] * 3 + [tone[i:i + 1024] for i in range(0, sr, 1024)] + [quiet] * 40:
            t += 0.02
            s = seg.feed(b, t)
            if s:
                out.append(s)
        self.assertEqual(len(out), 1)
        self.assertAlmostEqual(len(out[0].audio) / 16000, 1.5, delta=0.3)
        self.assertEqual(len(to_16k(tone, sr)), len(np.arange(0, sr, 3)))

    def test_mom_end_to_end(self):
        srv = HTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        with tempfile.TemporaryDirectory() as d:
            cfg = dict(DEFAULTS, ollama_url=f"http://127.0.0.1:{srv.server_port}/api/generate",
                       _moms=Path(d), _transcripts=Path(d))
            long = "\n".join(f"[10:00:{i % 60:02d}] Meeting: line {i} " + "word " * 20 for i in range(200))
            SEEN.clear()
            out = mg.generate_mom(long, datetime(2026, 10, 14, 14, 30), datetime(2026, 10, 14, 15, 30), cfg)
            self.assertGreater(len(SEEN), 2)  # chunked map step + final
            self.assertEqual(SEEN[-1]["model"], "llama3")
            doc = Document(str(out))
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("Q4 release plan", text)
            self.assertNotIn("Toyota", text)
            self.assertIn("Release Scope:", text)
            self.assertEqual(sum(1 for p in doc.paragraphs if p.style.name == "List Paragraph" and p.text), 3)
            t0 = {r.cells[0].text: r.cells[1].text for r in doc.tables[0].rows}
            self.assertEqual(t0["Date"], "Oct 14, 2026")
            self.assertEqual(t0["Time"], "02:30 to 03:30 PM")
            self.assertEqual(t0["Acme Participants"], "Jane Doe, John Roe")
            self.assertEqual(t0["Mahat Participants"], "Kishore")
            self.assertEqual(len(doc.tables[1].rows), 4)
            self.assertEqual(doc.tables[1].rows[1].cells[3].text, "20 Oct 2026")
        srv.shutdown()

    def test_transcript_writer(self):
        with tempfile.TemporaryDirectory() as d:
            w = TranscriptWriter(Path(d), datetime(2026, 10, 14, 14, 30))
            w.add(datetime(2026, 10, 14, 14, 31), "You", "Hello")
            w.finalize(datetime(2026, 10, 14, 15, 0))
            self.assertIn("Hello", "\n".join(p.text for p in Document(str(w.docx_path)).paragraphs))
            self.assertIn("[14:31:00] You: Hello", w.text())


if __name__ == "__main__":
    unittest.main()

"""Live transcript: appended to a Word document (and a .txt safety copy) as speech is recognised."""
import logging
import os
import threading
from datetime import datetime

from docx import Document
from docx.shared import Pt

log = logging.getLogger(__name__)


class TranscriptWriter:
    def __init__(self, folder, started: datetime):
        stamp = started.strftime("%Y-%m-%d_%H%M")
        self.docx_path = folder / f"Transcript_{stamp}.docx"
        self.txt_path = folder / f"Transcript_{stamp}.txt"
        self.lines = []
        self._lock = threading.Lock()
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = "Segoe UI", Pt(10)
        self.doc.add_heading("Meeting Transcript", level=1)
        self.doc.add_paragraph(f"Recording started: {started.strftime('%b %d, %Y %I:%M %p')}")
        self._save()

    def add(self, ts: datetime, label: str, text: str):
        line = f"[{ts.strftime('%H:%M:%S')}] {label}: {text}"
        with self._lock:
            self.lines.append(line)
            with open(self.txt_path, "a", encoding="utf-8") as f:
                f.write(line + "\n")
            p = self.doc.add_paragraph()
            r = p.add_run(f"[{ts.strftime('%H:%M:%S')}] ")
            r2 = p.add_run(f"{label}: ")
            r2.bold = True
            p.add_run(text)
            self._save()

    def _save(self):
        tmp = self.docx_path.with_suffix(".tmp")
        try:
            self.doc.save(tmp)
            os.replace(tmp, self.docx_path)  # fails while the file is open in Word -> retried next line
        except PermissionError:
            log.info("Transcript .docx is open in Word; will update on next segment")
        except Exception:
            log.exception("saving transcript docx")

    def finalize(self, ended: datetime):
        with self._lock:
            self.doc.add_paragraph(f"Recording ended: {ended.strftime('%b %d, %Y %I:%M %p')}")
            self._save()

    def text(self) -> str:
        with self._lock:
            return "\n".join(self.lines)

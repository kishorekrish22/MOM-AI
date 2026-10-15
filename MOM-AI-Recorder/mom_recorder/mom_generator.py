"""Transcript -> MOM (.docx) using the organisation's MOM template and the cloud-VM LLM."""
import copy
import json
import logging
import re
from datetime import datetime
from pathlib import Path

import requests
from docx import Document

from .config import resource_path

log = logging.getLogger(__name__)

CHUNK_CHARS = 6000
SINGLE_PASS_MAX = 9000

STYLE_EXAMPLE = {
    "overview": "The meeting focused on demonstrating the Mahat AI-based development platform and "
                "aligning Toyota's existing requirement, approval, testing, infrastructure, and "
                "deployment processes with the Mahat development approach.",
    "scope": "To establish a clear methodology for requirement sign-off, traceability, automated "
             "testing, NFR finalization, and hosting/deployment, so that MSRA can proceed with "
             "minimum rework and become a successful pilot for future Toyota applications.",
    "discussion_points": [
        {"title": "Requirement Sign-off & Traceability",
         "details": "Toyota emphasized that the UR/DUR must be clearly finalized and signed off before "
                    "development progresses. Mahat needs to demonstrate end-to-end traceability from "
                    "requirement to DUR/User Story to Code to Test Case to Test Result."},
        {"title": "Hosting, Repository & CI/CD",
         "details": "The solution is not SaaS-only. It can be deployed on-premise or on cloud based on "
                    "the finalized NFR. Both VM/on-prem and cloud options need to be evaluated early."}],
    "action_items": [
        {"activity": "Complete and obtain sign-off for the NFR.", "organization": "TKM",
         "owner": "Karthick Raja San", "due_by": "", "comments": ""},
        {"activity": "Schedule the detailed follow-up session", "organization": "Mahat",
         "owner": "Kishore", "due_by": "27 Oct 2026", "comments": ""}],
}

FINAL_PROMPT = """You are the minute-taker for {own_org}. Write the Minutes of the Meeting (MOM) from the meeting material below.

Rules:
- Use ONLY facts in the material. Never invent names, dates, decisions or numbers. If something is not stated, use null (or "" for text fields).
- Speaker labels "{mic}" and "{sys}" are audio sources, not people's names.
- Style: concise, professional, third person, past tense, like the example. Each discussion point has a short title and a 2-4 sentence paragraph. 5-10 discussion points at most.
- Action items: only real commitments or follow-ups; owner/organization only if stated; due_by only if a date was stated.
- Reply with ONE JSON object only, no markdown, with exactly these keys:
  overview (string), scope (string), mode (string or null: e.g. "Online - Microsoft Teams" / "Offline - In Person Meeting" only if clear),
  external_organization (string or null), external_participants (array of names), internal_participants (array of {own_org} names),
  discussion_points (array of {{"title","details"}}), action_items (array of {{"activity","organization","owner","due_by","comments"}}).

Example of the expected style and shape (different meeting):
{example}

MEETING MATERIAL:
{material}
"""

CHUNK_PROMPT = """Below is part {i} of {n} of a meeting transcript ("{mic}" and "{sys}" are audio sources, not names).
Write compact bullet-point notes: topics discussed, decisions, requirements, concerns, names mentioned
(with their organization if stated), dates, and any action items with owner/due date. Use only what is said.

TRANSCRIPT PART:
{text}
"""


class LLMError(RuntimeError):
    pass


class OllamaClient:
    def __init__(self, cfg):
        self.url, self.cfg = cfg["ollama_url"], cfg
        self.model = cfg["llm_model"]

    def _pick_model(self):
        if self.model:
            return self.model
        tags = self.url.rsplit("/api/", 1)[0] + "/api/tags"
        try:
            r = requests.get(tags, timeout=20)
            r.raise_for_status()
            models = r.json().get("models", [])
        except Exception as e:
            raise LLMError(f"Cannot reach the AI server at {tags} ({e}). Check the VM is running, the "
                           f"firewall allows port 11434 from this PC, or open that URL in a browser. "
                           f"You can also set \"llm_model\" in config.json.") from e
        if not models:
            raise LLMError("The AI server answered but lists no models. Set \"llm_model\" in config.json "
                           "(run `ollama list` on the VM to see the names).")
        self.model = models[0]["name"]
        log.info("Using LLM model: %s", self.model)
        return self.model

    def generate(self, prompt: str, as_json=False) -> str:
        body = {"model": self._pick_model(), "prompt": prompt, "stream": False,
                "options": {"temperature": 0.2, "num_ctx": 8192}}
        if as_json:
            body["format"] = "json"
        try:
            r = requests.post(self.url, json=body, timeout=self.cfg["llm_timeout_s"])
            r.raise_for_status()
            return r.json()["response"]
        except requests.RequestException as e:
            raise LLMError(f"LLM request failed: {e}") from e


def parse_json(text: str) -> dict:
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.M).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        a, b = text.find("{"), text.rfind("}")
        if a < 0 or b <= a:
            raise
        return json.loads(text[a:b + 1])


def chunk_transcript(text: str, size=CHUNK_CHARS):
    chunks, cur = [], ""
    for line in text.splitlines():
        if cur and len(cur) + len(line) > size:
            chunks.append(cur)
            cur = ""
        cur += line + "\n"
    if cur.strip():
        chunks.append(cur)
    return chunks


def extract_mom_data(transcript: str, cfg: dict, llm=None) -> dict:
    llm = llm or OllamaClient(cfg)
    common = {"own_org": cfg["own_organization"], "mic": cfg["mic_label"], "sys": cfg["system_label"]}
    if len(transcript) <= SINGLE_PASS_MAX:
        material = transcript
    else:
        chunks = chunk_transcript(transcript)
        notes = [llm.generate(CHUNK_PROMPT.format(i=i, n=len(chunks), text=c, **common))
                 for i, c in enumerate(chunks, 1)]
        material = "NOTES FROM THE FULL TRANSCRIPT:\n\n" + "\n\n".join(notes)
    prompt = FINAL_PROMPT.format(example=json.dumps(STYLE_EXAMPLE, indent=1), material=material, **common)
    raw = llm.generate(prompt, as_json=True)
    try:
        return parse_json(raw)
    except json.JSONDecodeError:
        log.warning("LLM returned invalid JSON, retrying once")
        raw = llm.generate(prompt + "\n\nYour previous reply was not valid JSON. Reply with the JSON object only.",
                           as_json=True)
        return parse_json(raw)


# ---------------------------------------------------------------- docx filling

def _str(v) -> str:
    return "" if v is None else str(v).strip()


def _names(v) -> str:
    return ", ".join(_str(x) for x in v if _str(x)) if isinstance(v, list) else _str(v)


def _set_par_text(p, text):
    """Replace paragraph text, keeping the first run's formatting."""
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        r = p.add_run(text)
        rpr = p._p.find("./{*}pPr/{*}rPr")  # paragraph-mark formatting carries the template font
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))
    for el in p._p.findall("{*}proofErr"):
        p._p.remove(el)


def _set_cell(cell, text):
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    _set_par_text(cell.paragraphs[0], text)


def _is_bold(run):
    return run.bold is True


def fill_template(data: dict, started: datetime, ended: datetime, cfg: dict, out: Path) -> Path:
    doc = Document(str(resource_path("templates/MOM_Template.docx")))
    body = doc.element.body
    paras = doc.paragraphs
    t_info, t_act = doc.tables[0], doc.tables[1]

    # Overview / scope (paragraph after each heading)
    for i, p in enumerate(paras):
        h = p.text.replace("\xa0", "").strip()
        if h == "Overview":
            _set_par_text(paras[i + 1], _str(data.get("overview")) or "Not stated.")
        elif h == "Scope of the Discussion":
            _set_par_text(paras[i + 1], _str(data.get("scope")) or "Not stated.")

    # Details table
    ext = _str(data.get("external_organization"))
    mode = _str(data.get("mode")) or cfg["default_mode"]
    start_t, end_t = started.strftime("%I:%M"), ended.strftime("%I:%M %p")
    values = {
        "Date": started.strftime("%b %d, %Y"),
        "Time": f"{start_t} to {end_t}",
        "Mode": mode,
    }
    for row in t_info.rows[1:]:
        label = row.cells[0].text.strip()
        if label in values:
            _set_cell(row.cells[1], values[label])
        elif label.endswith("Participants"):
            own = cfg["own_organization"]
            if label.startswith(own):
                _set_cell(row.cells[1], _names(data.get("internal_participants")) or "Not captured")
            else:
                if ext:
                    _set_cell(row.cells[0], f"{ext} Participants")
                _set_cell(row.cells[1], _names(data.get("external_participants")) or "Not captured")

    # Discussion points: clone the template's numbered paragraph
    items = [i for i in (data.get("discussion_points") or []) if isinstance(i, dict) and _str(i.get("title"))]
    protos = [p for p in paras if p.style.name == "List Paragraph" and p.text.strip()]
    if protos and items:
        proto = protos[0]
        runs = proto.runs
        title_r, br_r = runs[0], runs[1]
        body_r = next(r for r in runs[2:] if not _is_bold(r) and r.text)
        kids = list(proto._p)
        keep = {kids.index(x._r) for x in (title_r, br_r, body_r)}
        for it in items:
            new = copy.deepcopy(proto._p)
            for idx, ch in enumerate(list(new)):
                if not ch.tag.endswith("}pPr") and idx not in keep:
                    new.remove(ch)
            ts = [t for t in new.iter() if t.tag.endswith("}t")]
            ts[0].text = _str(it["title"]).rstrip(":") + ":"
            ts[-1].text = _str(it.get("details"))
            proto._p.addprevious(new)
    for p in protos:  # never leave the template's sample text behind
        body.remove(p._p)

    # Action items: clone data rows alternately for banding
    acts = [a for a in (data.get("action_items") or []) if isinstance(a, dict) and _str(a.get("activity"))]
    rows = t_act.rows
    protos_tr = [rows[1]._tr, rows[2]._tr if len(rows) > 2 else rows[1]._tr]
    protos_tr = [copy.deepcopy(x) for x in protos_tr]
    for r in rows[1:]:
        t_act._tbl.remove(r._tr)
    if not acts:
        acts = [{"activity": "No action items were identified.", "organization": "", "owner": "",
                 "due_by": "", "comments": ""}]
    from docx.table import _Row
    for n, a in enumerate(acts):
        tr = copy.deepcopy(protos_tr[n % 2])
        t_act._tbl.append(tr)
        row = _Row(tr, t_act)
        for cell, key in zip(row.cells, ("activity", "organization", "owner", "due_by", "comments")):
            _set_cell(cell, _str(a.get(key)))

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def generate_mom(transcript: str, started: datetime, ended: datetime, cfg: dict, llm=None) -> Path:
    if not transcript.strip():
        raise LLMError("Transcript is empty - nothing to summarise.")
    data = extract_mom_data(transcript, cfg, llm)
    out = cfg["_moms"] / f"MOM_{started.strftime('%Y-%m-%d_%H%M')}.docx"
    (cfg["_moms"] / f"{out.stem}.json").write_text(json.dumps(data, indent=2, ensure_ascii=False),
                                                     encoding="utf-8")
    return fill_template(data, started, ended, cfg, out)
